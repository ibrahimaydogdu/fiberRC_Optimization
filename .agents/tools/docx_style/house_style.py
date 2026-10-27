"""Proje Word biçim kuralları (kullanıcı kuralı, 2026-10-17).

Ana Başlık (Title)  : Times New Roman 14 pt, ortalı, siyah, kalın, 12/12 pt aralık
Başlık 1            : TNR 14 pt, kalın, siyah, 12/12
Başlık 2            : TNR 13 pt, kalın, siyah, 6/6
Başlık 3 (4+ dahil) : TNR 12 pt, kalın, siyah, 6/6
Normal              : TNR 12 pt, siyah, 6/6
Tablo içi (Compact) : TNR 12 pt, siyah, aralık 0
Sayfa               : A4, kenar boşluğu 2,5 cm (kullanıcı belirtmedi; varsayılan)

Kullanım:
  - pandoc için: build_reference(path) -> reference.docx
  - python-docx ile üretilen belgeler için: apply_house_style(doc)
"""
import subprocess
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"
C, L = WD_ALIGN_PARAGRAPH.CENTER, None

# stil adı: (punto, kalın, hizalama, önce, sonra)
SPEC = {
    "Title": (14, True, C, 12, 12),
    "Heading 1": (14, True, L, 12, 12),
    "Heading 2": (13, True, L, 6, 6),
    "Heading 3": (12, True, L, 6, 6),
    "Heading 4": (12, True, L, 6, 6),
    "Heading 5": (12, True, L, 6, 6),
    "Heading 6": (12, True, L, 6, 6),
    "Normal": (12, False, L, 6, 6),
    "Body Text": (12, False, L, 6, 6),
    "First Paragraph": (12, False, L, 6, 6),
    "Block Text": (12, False, L, 6, 6),
    "Compact": (12, False, L, 0, 0),
    "Table Caption": (12, False, L, 6, 6),
    "Image Caption": (12, False, C, 6, 6),
    "Caption": (12, False, L, 6, 6),
}


def _fix_fonts(style, bold):
    f = style.font
    f.name = FONT
    f.size = None  # boyut aşağıda atanır
    f.bold = bold
    f.italic = False if style.name.startswith(("Heading", "Title")) else f.italic
    f.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(a), FONT)
    color = rpr.find(qn("w:color"))
    if color is not None:
        for a in ("w:themeColor", "w:themeShade", "w:themeTint"):
            if color.get(qn(a)) is not None:
                del color.attrib[qn(a)]
        color.set(qn("w:val"), "000000")


def style_styles(doc):
    for name, (size, bold, align, before, after) in SPEC.items():
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        _fix_fonts(st, bold)
        st.font.size = Pt(size)
        pf = st.paragraph_format
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        pf.line_spacing = 1.0
        if align is not None:
            pf.alignment = align


def set_page(doc):
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.5)


def style_table_cells(doc):
    """Tablo içi: TNR 12, siyah, aralık yok (python-docx belgeleri için açık biçim)."""
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
                    for r in p.runs:
                        r.font.name = FONT
                        r.font.size = Pt(12)
                        r.font.color.rgb = RGBColor(0, 0, 0)
                        rpr = r._r.get_or_add_rPr()
                        rf = rpr.find(qn("w:rFonts"))
                        if rf is not None:
                            rf.set(qn("w:eastAsia"), FONT)


def apply_house_style(doc):
    style_styles(doc)
    set_page(doc)
    style_table_cells(doc)
    return doc


import os
import shutil


def _find_pandoc():
    p = shutil.which("pandoc")
    if p:
        return p
    candidates = [
        os.path.expanduser(r"~\AppData\Local\Programs\pandoc\pandoc.exe"),
        r"C:\Program Files\Pandoc\pandoc.exe",
        r"C:\Program Files (x86)\Pandoc\pandoc.exe",
    ]
    for c in candidates:
        if os.path.isfile(c):
            return c
    return "pandoc"


def build_reference(path):
    """pandoc varsayılan reference.docx'ini alıp kurallara göre biçimler."""
    pandoc_bin = _find_pandoc()
    data = subprocess.run(
        [pandoc_bin, "--print-default-data-file", "reference.docx"],
        check=True, capture_output=True).stdout
    with open(path, "wb") as fh:
        fh.write(data)
    doc = Document(path)
    style_styles(doc)
    set_page(doc)
    doc.save(path)
    return path


if __name__ == "__main__":
    import sys
    print(build_reference(sys.argv[1]))
