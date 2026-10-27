"""Markdown -> Word dönüştürücü (proje biçim kurallarıyla).

Kullanım:  python -I -X utf8 md2docx.py girdi.md cikti.docx [ek pandoc argümanları]
Adımlar : reference.docx üret (yoksa) -> pandoc + Lua süzgeci -> denklem tablolarını biçimle.
Kurallar: house_style.py ve captions_equations.lua.
"""
import os
import re
import shutil
import subprocess
import sys
import tempfile
import warnings

warnings.filterwarnings("ignore", message="style lookup by style_id")

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
import house_style  # noqa: E402

M = "http://schemas.openxmlformats.org/officeDocument/2006/math"


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _find_pandoc():
    p = shutil.which("pandoc")
    if p:
        return p
    candidates = [
        os.path.expanduser(r"~\AppData\Local\Programs\pandoc\pandoc.exe"),
        r"C:\Program Files\Pandoc\pandoc.exe",
        r"C:\Program Files (x86)\Pandoc\pandoc.exe",
    ]
    for c in candidates:
        if os.path.isfile(c):
            return c
    return "pandoc"


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), str(v))
    return e


def _is_equation_table(tbl):
    if len(tbl.rows) != 1 or len(tbl.columns) != 2:
        return False
    tc0 = tbl.rows[0].cells[0]._tc
    tc1 = tbl.rows[0].cells[1]._tc

    # 1. Hücre 0'da matematik öğesi (oMath veya oMathPara) var mı?
    has_math = (len(tc0.findall("{%s}oMath" % M)) > 0 or
                len(tc0.findall("{%s}oMathPara" % M)) > 0)

    # 2. Hücre 1'de SEQ Denklem alanı var mı?
    tc1_xml = tc1.xml
    has_seq = "SEQ Denklem" in tc1_xml or ("Denklem" in tc1_xml and "Figur" not in tc1_xml)

    # 3. Hücre 1'de parantezli denklem numarası kalıbı var mı?
    texts = [t.text for t in tc1.findall("{%s}t" % W_NS) if t.text]
    full_text = "".join(texts).strip()
    cell_text = tbl.rows[0].cells[1].text.strip()
    has_num_pattern = (bool(re.fullmatch(r"\(\s*\d*\s*\)", full_text)) or
                       bool(re.fullmatch(r"\(\s*\d*\s*\)", cell_text)))

    return has_math or (has_seq and has_num_pattern) or (has_num_pattern and has_seq)


def _format_equation_table(tbl):
    tblPr = tbl._tbl.tblPr
    # tblStyle ve diğer kenarlık/düzen özelliklerini temizle
    for tag in ("w:tblStyle", "w:tblBorders", "w:tblCellMar", "w:tblLayout", "w:tblW", "w:tblLook"):
        for old in tblPr.findall(qn(tag)):
            tblPr.remove(old)

    tblPr.append(_el("w:tblW", w=5000, type="pct"))

    # Tablo düzeyinde tüm kenarlıkları sıfırla
    borders = _el("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(_el("w:" + side, val="none", sz=0, space=0, color="auto"))
    tblPr.append(borders)

    tblPr.append(_el("w:tblLayout", type="fixed"))

    mar = _el("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        mar.append(_el("w:" + side, w=0, type="dxa"))
    tblPr.append(mar)

    # tblLook sıfırla; böylece firstRow/lastRow stil kuralları (alt kenarlık vb.) tetiklenmez
    tblPr.append(_el("w:tblLook", val="0000", firstRow=0, lastRow=0, firstColumn=0, lastColumn=0, noHBand=0, noVBand=0))

    # Satır ve hücre seviyesinde kenarlıkları ve koşullu biçimlendirmeyi temizle
    for row in tbl.rows:
        trPr = row._tr.get_or_add_trPr()
        for cnf in trPr.findall(qn("w:cnfStyle")):
            trPr.remove(cnf)

        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            # Hücredeki mevcut tcBorders'ı kaldırıp kesin "none" kenarlıklar ata
            for old in tcPr.findall(qn("w:tcBorders")):
                tcPr.remove(old)
            tcBorders = _el("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                tcBorders.append(_el("w:" + side, val="none", sz=0, space=0, color="auto"))
            tcPr.append(tcBorders)

            # Paragraf seviyesinde pBdr kenarlık varsa temizle
            for p in cell.paragraphs:
                pPr = p._p.get_or_add_pPr()
                for old in pPr.findall(qn("w:pBdr")):
                    pPr.remove(old)

    # denklem sola hizalı
    for para in tbl._tbl.iter("{%s}oMathPara" % M):
        pr = para.find("{%s}oMathParaPr" % M)
        if pr is None:
            pr = OxmlElement("m:oMathParaPr")
            para.insert(0, pr)
        for jc in pr.findall("{%s}jc" % M):
            pr.remove(jc)
        jc = OxmlElement("m:jc")
        jc.set("{%s}val" % M, "left")
        pr.append(jc)


def postprocess(path):
    doc = Document(path)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if p.style is None or p.style.name != "Compact":
                        p.style = doc.styles["Compact"]
        if _is_equation_table(tbl):
            _format_equation_table(tbl)
    doc.save(path)


def convert(src, dst, extra=()):
    ref = os.path.join(HERE, "reference.docx")
    if not os.path.exists(ref):
        house_style.build_reference(ref)
    pandoc_bin = _find_pandoc()
    # Windows'ta pandoc, ASCII dışı yollardaki (ör. "Çalışmaları") dosyaları açamıyor:
    # reference.docx ve Lua süzgeci geçici ASCII dizine kopyalanır.
    with tempfile.TemporaryDirectory() as tmp:
        ref_t = shutil.copy(ref, os.path.join(tmp, "reference.docx"))
        lua_t = shutil.copy(os.path.join(HERE, "captions_equations.lua"),
                            os.path.join(tmp, "captions_equations.lua"))
        cmd = [pandoc_bin, src, "-o", dst, "--from", "markdown+tex_math_dollars",
               "--reference-doc", ref_t, "--lua-filter", lua_t,
               "--resource-path", os.path.dirname(os.path.abspath(src))] + list(extra)
        subprocess.run(cmd, check=True)
    postprocess(dst)
    return dst


if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    print(convert(sys.argv[1], sys.argv[2], sys.argv[3:]))
